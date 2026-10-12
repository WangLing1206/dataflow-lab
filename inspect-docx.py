import sys
import docx

path = sys.argv[1] if len(sys.argv) > 1 else r'e:\2026-2027-1\web6\docs\DataFlow-Lab-网站关键亮点说明.docx'
doc = docx.Document(path)
sec = doc.sections[0]
print('页面: %.1f x %.1f cm, 边距 %.1f cm' % (sec.page_width.cm, sec.page_height.cm, sec.left_margin.cm))
print('段落数:', len(doc.paragraphs), '表格数:', len(doc.tables),
      '图片数:', sum(1 for p in doc.paragraphs if 'graphic' in p._p.xml))
print('--- 结构（style / 字号 / 对齐 / 文本） ---')
for p in doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    sz = p.runs[0].font.size.pt if (p.runs and p.runs[0].font.size) else None
    al = str(p.alignment).replace('WD_ALIGN_PARAGRAPH.', '') if p.alignment is not None else '-'
    sty = p.style.name
    mark = '  ' if sty == 'Normal' else '* '
    print(f'{mark}{sty:<10} {str(sz):<5} {al:<8} {t[:58]}')
